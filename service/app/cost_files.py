"""Reading a seller's cost file, and matching its rows to the variants TikTok returned.

CST-2. This module is pure: it takes the file's bytes and the shop's variants and returns
rows and outcomes. It does not know where the file came from, because how a file reaches
the store is not settled (see CLAUDE.md and A10.8). The four upload operations are not
served until it is. `tests/test_cost_files.py` runs every rule below.

WHAT IS READ

CSV and Excel `.xlsx` only, as the owner ruled (cost files are Excel and CSV only, because
Word and PDF carry no columns). The first non-empty row is the header row. For Excel only
the first worksheet is read, and formulas are read as their last saved values. XML inside
an `.xlsx` is parsed through `defusedxml`, which `openpyxl` uses when it is installed.

MONEY

A cost is read through `Decimal`, never a binary float (A13 rule 3, A29.13). Excel stores a
typed 3.40 as a binary double, so a numeric cell is converted through its shortest decimal
representation, which gives back what the seller typed for any amount of two decimals. A
pound sign, spaces and thousands commas are accepted. More than two decimals, a negative
amount, or text that is not a number makes the row unmatched with a reason, rather than
rounding a cost the seller did not write.

MATCHING, AS THE CONTRACT STATES IT

- Matching never creates a variant. A row naming one TikTok has not returned is unmatched.
- With `match_on = seller_sku`, the key is compared with each variant's seller SKU. Where
  that finds nothing, it is compared with TikTok's own SKU id, because the contract says
  matching falls back to it. `matched_on` says which one matched.
- With `match_on = tiktok_sku_id`, only TikTok's SKU id is compared.

**Derived, not ruled.** Keys are compared after trimming surrounding spaces, and case is
kept, because nothing says TikTok treats seller SKUs as case-insensitive. Where two rows of
one file resolve to the same variant, the first is matched and each later one is a
duplicate naming the earlier row, because applying both would leave the file's own order
deciding which cost wins without the seller seeing it.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import Any
from uuid import UUID

CSV_TYPE = "text/csv"
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_ROWS = 20000


class FileUnreadable(ValueError):
    """The file cannot be read as a cost file. The message is shown to the seller."""


@dataclass
class ParsedFile:
    columns: list[str]
    rows: list[dict[str, Any]]


@dataclass
class Variant:
    sku_id: UUID
    seller_sku: str | None
    tiktok_sku_id: str


@dataclass
class RowOutcome:
    row_number: int
    outcome: str  # matched, unmatched or duplicate
    matched_on: str | None = None
    sku_id: UUID | None = None
    seller_sku: str | None = None
    unit_cost_minor: int | None = None
    packing_minor: int | None = None
    postage_minor: int | None = None
    reason: str | None = None
    raw: dict[str, Any] = field(default_factory=dict)


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        # repr gives the shortest string that round-trips, which is what was typed.
        return repr(value)
    return str(value).strip()


def parse(data: bytes, content_type: str) -> ParsedFile:
    if content_type == CSV_TYPE:
        table = _read_csv(data)
    elif content_type == XLSX_TYPE:
        table = _read_xlsx(data)
    else:
        raise FileUnreadable("Only CSV and Excel (.xlsx) files can be read.")

    table = [r for r in table if any(c != "" for c in r)]
    if table:
        table[0] = [_cell_text(h) for h in table[0]]
    if not table:
        raise FileUnreadable("The file has no rows.")
    header = table[0]
    if any(h == "" for h in header):
        raise FileUnreadable("Every column in the first row needs a heading.")
    if len(set(header)) != len(header):
        raise FileUnreadable("Two columns in the first row have the same heading.")
    body = table[1:]
    if len(body) > MAX_ROWS:
        raise FileUnreadable(f"The file has more than {MAX_ROWS:,} rows. Split it into smaller files and upload each one.")
    rows = [
        {h: (r[i] if i < len(r) else "") for i, h in enumerate(header)}
        for r in body
    ]
    return ParsedFile(columns=header, rows=rows)


def _read_csv(data: bytes) -> list[list[str]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            # Excel on Windows saves CSV in Windows-1252 unless told otherwise.
            text = data.decode("cp1252")
        except UnicodeDecodeError as exc:
            raise FileUnreadable("MyShopEdge could not read the text in this CSV file. Save it again as CSV UTF-8 and upload it.") from exc
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text))]


def _read_xlsx(data: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise FileUnreadable("The Excel file could not be opened.") from exc
    try:
        sheet = wb.worksheets[0]
        # Values keep their type here. A number in a key column has to be recognised as a
        # number later, because Excel cannot hold a TikTok SKU id as one.
        return [
            ["" if c is None else (c.strip() if isinstance(c, str) else c) for c in row]
            for row in sheet.iter_rows(values_only=True)
        ]
    finally:
        wb.close()


_AMOUNT = re.compile(r"^£?\s*(\d{1,3}(,\d{3})+|\d+)(\.\d+)?$")


def read_amount(text: str) -> int:
    """Pounds as written by a seller, to integer pence. Raises ValueError with a reason."""
    t = text.strip()
    if t == "":
        raise ValueError("is empty")
    if t.startswith("-"):
        raise ValueError("is negative")
    if not _AMOUNT.match(t):
        raise ValueError(f'is "{text}", which is not an amount')
    try:
        d = Decimal(t.replace("£", "").replace(",", "").strip())
    except InvalidOperation as exc:
        raise ValueError(f'is "{text}", which is not an amount') from exc
    if d != d.quantize(Decimal("0.01")):
        raise ValueError(f'is "{text}", which has more than two decimal places')
    return int(d * 100)


_KEY_WORDS = {
    "tiktok_sku_id": ("tiktok sku id", "sku id", "tiktok sku"),
    "seller_sku": ("seller sku", "sku", "seller sku code"),
}
_COST_WORDS = ("unit cost", "cost price", "product cost", "cost")
_PACKING_WORDS = ("packing", "packaging", "packing cost")
_POSTAGE_WORDS = ("postage", "shipping", "delivery", "postage cost")


def _norm(h: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", h.lower()).strip()


def suggest_mapping(columns: list[str]) -> dict[str, Any] | None:
    """A proposal from the headings, never applied without the seller confirming it.

    Returns None when no heading plainly names both a key and a cost, rather than guessing.
    """
    norm = {_norm(c): c for c in columns}

    def find(words: tuple[str, ...]) -> str | None:
        for w in words:
            if w in norm:
                return norm[w]
        return None

    match_on, key = None, None
    for kind in ("tiktok_sku_id", "seller_sku"):
        found = find(_KEY_WORDS[kind])
        if found:
            match_on, key = kind, found
            break
    cost = find(_COST_WORDS)
    if not (key and cost):
        return None
    return {
        "match_on": match_on,
        "key_column": key,
        "cost_column": cost,
        "packing_column": find(_PACKING_WORDS),
        "postage_column": find(_POSTAGE_WORDS),
        "currency": "GBP",
    }


def match(parsed: ParsedFile, mapping: dict[str, Any], variants: list[Variant]) -> list[RowOutcome]:
    for col in ("key_column", "cost_column", "packing_column", "postage_column"):
        name = mapping.get(col)
        if name and name not in parsed.columns:
            raise FileUnreadable(f'The file has no column headed "{name}".')

    by_seller = {}
    for v in variants:
        if v.seller_sku and v.seller_sku.strip():
            by_seller.setdefault(v.seller_sku.strip(), v)
    by_tiktok = {v.tiktok_sku_id.strip(): v for v in variants}

    seen: dict[UUID, int] = {}
    out: list[RowOutcome] = []
    for i, row in enumerate(parsed.rows, start=2):  # row 1 is the headings
        cell = row.get(mapping["key_column"], "")
        o = RowOutcome(row_number=i, outcome="unmatched",
                       raw={k: _cell_text(v) for k, v in row.items()})
        if isinstance(cell, float) or (isinstance(cell, int) and abs(cell) >= 10**15):
            # Found on 24 September by writing a real .xlsx: a 19 digit TikTok SKU id typed
            # into a number cell reads back as 1.7291e+18. Excel keeps 15 significant digits,
            # so the code was changed before the file was saved and cannot be recovered.
            o.reason = (
                "This SKU was saved as a number, and Excel keeps only 15 digits, so the code "
                "was changed. Format the column as text, retype the codes and upload again."
            )
            out.append(o)
            continue
        key = _cell_text(cell)

        if key == "":
            o.reason = "The row has no SKU, so it cannot be matched."
            out.append(o)
            continue

        variant, on = None, None
        if mapping["match_on"] == "seller_sku" and key in by_seller:
            variant, on = by_seller[key], "seller_sku"
        elif key in by_tiktok:
            variant, on = by_tiktok[key], "tiktok_sku_id"
        if variant is None:
            o.seller_sku = key if mapping["match_on"] == "seller_sku" else None
            o.reason = (
                f'"{key}" is not a SKU TikTok has returned for this shop. '
                "Nothing is created for it. Check the code, or add the cost by hand."
            )
            out.append(o)
            continue

        # The column being read, so a bad packing or postage cell is not reported as a bad
        # cost (copy audit, 8 October 2026).
        reading = mapping["cost_column"]
        try:
            o.unit_cost_minor = read_amount(_cell_text(row.get(mapping["cost_column"], "")))
            for col, attr in (("packing_column", "packing_minor"), ("postage_column", "postage_minor")):
                if mapping.get(col) and _cell_text(row.get(mapping[col], "")) != "":
                    reading = mapping[col]
                    setattr(o, attr, read_amount(_cell_text(row[mapping[col]])))
        except ValueError as err:
            o.sku_id, o.matched_on, o.seller_sku = variant.sku_id, on, variant.seller_sku
            o.reason = f'The amount in the "{reading}" column {err}, so this row was not read.'
            o.unit_cost_minor = None
            out.append(o)
            continue

        o.sku_id, o.matched_on, o.seller_sku = variant.sku_id, on, variant.seller_sku
        if variant.sku_id in seen:
            o.outcome = "duplicate"
            o.reason = f"This variant already appears on row {seen[variant.sku_id]} of this file."
        else:
            o.outcome = "matched"
            seen[variant.sku_id] = i
        out.append(o)
    return out
