"""Runs the scheduled exports (MON-8) against a local PostgreSQL and a local stand-in for S3.

    pip install "moto[server]"
    DATABASE_URL=postgresql://mse_migrator@/mse_sched_1?host=/tmp/pgcheck&port=5499 \\
        python3 testdata/scheduled_exports_check.py

Written 9 October 2026, when scheduled exports were built. The database must be built by
`service/scripts/migrate.py` through 0030 and hold the year dataset
(`testdata/load_rows.py testdata/rows_year.json`), whose one account owns the shop below. The
URL must be a role that may set `mse_app`, as the service's own does. It starts `moto` on a
local port, as `storage_check.py` does, and reaches no network. It adds a second account of
its own, and removes every row it made, whether it passes or not.

What it proves, through the handlers and the runner themselves:

- create, list, pause, resume, change and delete; a replayed Idempotency-Key creates once;
  a weekly schedule on the cash basis is refused;
- another account's schedule answers 404 to a change and to a removal, and is not listed;
- a weekly and a monthly schedule run on their day and not on another, and a paused one
  does not run;
- with storage unconfigured the schedule is not marked as run, and no export or notice is
  written;
- a run writes one ready export, sets `last_run_at` and raises one notice, and a second run
  the same London day writes nothing;
- each file's totals equal the Money screen (`getMoney`) for the period the run covers.

**It proves the code path, not R2**, for the reason `storage_check.py` gives.
"""

import csv
import io
import json
import os
import socket
import subprocess
import sys
import time
import uuid
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path

import httpx

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "service"))

SHOP = uuid.UUID("8a773a13-73b5-a382-7dd0-fda02e950369")
ACCOUNT = uuid.UUID("56e487ea-e0fa-3691-7857-724855e716fc")

ok = bad = 0


def check(cond: bool, what: str) -> None:
    global ok, bad
    if cond:
        ok += 1
        print(f"  ok   {what}")
    else:
        bad += 1
        print(f"  FAIL {what}")


def free_port() -> int:
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


def at(day: str) -> datetime:
    """06:47 in London on a summer day, the hour the daily job runs."""
    return datetime.fromisoformat(f"{day}T05:47:00+00:00").astimezone(timezone.utc)


def pence(text) -> int | None:
    if text in (None, ""):
        return None
    try:
        return int(Decimal(str(text)) * 100)
    except ArithmeticError:
        return None


def totals_from_rows(rows: list[list]) -> dict[str, int | None]:
    """The month summary's Totals block, in pence, read back out of the file."""
    out = {}
    for r in rows:
        if len(r) >= 4 and r[0] == "Totals":
            out[r[1]] = pence(r[3])
    return out


def read_file(data: bytes, fmt: str) -> list[list]:
    if fmt == "csv":
        return list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(data)).active
    return [["" if v is None else v for v in row] for row in ws.iter_rows(values_only=True)]


def main() -> int:
    port = free_port()
    server = subprocess.Popen([sys.executable, "-m", "moto.server", "-p", str(port)],
                              stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    endpoint = f"http://127.0.0.1:{port}"
    os.environ["NO_PROXY"] = os.environ["no_proxy"] = "127.0.0.1"
    for _ in range(50):
        try:
            httpx.get(endpoint, timeout=0.5)
            break
        except httpx.HTTPError:
            time.sleep(0.2)
    # Every storage variable but the bucket, so the first run meets `storage_unconfigured`.
    os.environ.update(AWS_ENDPOINT_URL_S3=endpoint, AWS_REGION="auto",
                      AWS_ACCESS_KEY_ID="check", AWS_SECRET_ACCESS_KEY="check")
    os.environ.pop("S3_BUCKET", None)
    import boto3

    boto3.client("s3", region_name="us-east-1", endpoint_url=endpoint).create_bucket(Bucket="mse-sched-check")

    from fastapi import Response

    from app import storage
    from app.auth import Account
    from app.db import pool
    from app.export_schedules import (ScheduleChange, ScheduleIn, create_export_schedule,
                                      delete_export_schedule, list_export_schedules, period_for,
                                      run_due, update_export_schedule)
    from app.money_view import get_money
    from app.problems import Problem

    tag = uuid.uuid4().hex[:8]
    other, other_shop = uuid.uuid4(), uuid.uuid4()
    me = Account(id=ACCOUNT, email="a", name=None, subject="a")
    them = Account(id=other, email="b", name=None, subject="b")
    with pool().connection() as conn, conn.transaction():
        found = conn.execute("select account_id from shops where id = %s", (str(SHOP),)).fetchone()
        if not found or found[0] != ACCOUNT:
            print("The year dataset is not loaded: load testdata/rows_year.json first.")
            server.terminate()
            return 1
        conn.execute("insert into accounts (id, email, auth_subject) values (%s, %s, %s)",
                     (str(other), f"{other}@check.invalid", f"check-{tag}"))
        conn.execute("insert into shops (id, account_id, tiktok_shop_id, connection_status) "
                     "values (%s, %s, %s, 'connected')", (str(other_shop), str(other), f"check-{tag}"))

    def body(resp) -> dict:
        return json.loads(resp.body)

    def problem(fn) -> tuple[int, str] | None:
        try:
            fn()
        except Problem as p:
            return p.status_code, p.code
        return None

    def state(schedule_id):
        with pool().connection() as conn:
            last = conn.execute("select last_run_at from export_schedules where id = %s",
                                (schedule_id,)).fetchone()[0]
            exports = conn.execute(
                "select id, status, kind, format, basis, period_start, period_end, storage_key "
                "from exports where shop_id = %s order by created_at, id", (str(SHOP),)).fetchall()
            notices = conn.execute(
                "select entity_id, title, body, dedupe_key from notifications "
                "where account_id = %s and type = 'scheduled_export_ready' order by created_at",
                (str(ACCOUNT),)).fetchall()
        return last, exports, notices

    try:
        print("the four operations")
        weekly_in = ScheduleIn(kind="month_summary", format="csv", basis="sales", cadence="weekly", day_of_week=1)
        r1 = create_export_schedule(weekly_in, me, SHOP, f"sched-{tag}-1")
        r1b = create_export_schedule(weekly_in, me, SHOP, f"sched-{tag}-1")
        weekly = body(r1)
        check(r1.status_code == 201 and weekly["cadence"] == "weekly" and weekly["day_of_week"] == 1
              and weekly["active"] is True and weekly["last_run_at"] is None, "createExportSchedule answers 201 with the schedule")
        check(r1b.status_code == 201 and body(r1b)["id"] == weekly["id"], "a replayed Idempotency-Key returns the first schedule")
        monthly = body(create_export_schedule(
            ScheduleIn(kind="month_summary", format="xlsx", basis="cash", cadence="monthly", day_of_month=5), me, SHOP, None))
        spare = body(create_export_schedule(
            ScheduleIn(kind="ledger", format="csv", basis="sales", cadence="monthly", day_of_month=28), me, SHOP, None))
        listed = list_export_schedules(me, SHOP).schedules
        check([str(s.id) for s in listed] == [weekly["id"], monthly["id"], spare["id"]],
              f"listExportSchedules returns the three, oldest first, the replay adding none ({len(listed)})")
        check(problem(lambda: create_export_schedule(
            ScheduleIn(kind="ledger", format="csv", basis="cash", cadence="weekly", day_of_week=3), me, SHOP, None))
            == (422, "validation_failed"), "a weekly schedule on the cash basis is refused")
        check(problem(lambda: create_export_schedule(
            ScheduleIn(kind="ledger", format="csv", basis="sales", cadence="monthly"), me, SHOP, None))
            == (422, "validation_failed"), "a monthly schedule without its day is refused")

        paused = update_export_schedule(ScheduleChange(active=False), me, SHOP, uuid.UUID(spare["id"]))
        check(paused.active is False and paused.day_of_month == 28, "a pause changes only active")
        changed = update_export_schedule(ScheduleChange(cadence="weekly", day_of_week=4), me, SHOP, uuid.UUID(spare["id"]))
        check(changed.cadence == "weekly" and changed.day_of_week == 4 and changed.day_of_month is None
              and changed.active is False, "a change to weekly clears the day of the month and keeps the pause")
        check(problem(lambda: update_export_schedule(ScheduleChange(basis="cash"), me, SHOP, uuid.UUID(spare["id"])))
              == (422, "validation_failed"), "a change to the cash basis on a weekly schedule is refused")
        resumed = update_export_schedule(ScheduleChange(active=True, cadence="monthly", day_of_month=28),
                                         me, SHOP, uuid.UUID(spare["id"]))
        check(resumed.active and resumed.cadence == "monthly" and resumed.day_of_month == 28
              and resumed.day_of_week is None, "resume and a change back to monthly")

        print("another account")
        check(list_export_schedules(them, other_shop).schedules == [], "the other account lists nothing of this shop's")
        check(problem(lambda: update_export_schedule(ScheduleChange(active=False), them, other_shop, uuid.UUID(weekly["id"])))
              == (404, "schedule_not_found"), "the other account's change answers 404")
        check(problem(lambda: delete_export_schedule(them, other_shop, uuid.UUID(weekly["id"])))
              == (404, "schedule_not_found"), "the other account's removal answers 404")
        check(problem(lambda: update_export_schedule(ScheduleChange(active=False), me, SHOP, uuid.uuid4()))
              == (404, "schedule_not_found"), "a schedule that does not exist answers the same 404")
        check(len(list_export_schedules(me, SHOP).schedules) == 3, "the owner's three are untouched")

        deleted = delete_export_schedule(me, SHOP, uuid.UUID(spare["id"]))
        check(deleted.status_code == 204, "deleteExportSchedule answers 204")
        check(problem(lambda: delete_export_schedule(me, SHOP, uuid.UUID(spare["id"])))
              == (404, "schedule_not_found"), "a second removal answers 404")
        check(len(list_export_schedules(me, SHOP).schedules) == 2, "two are left")

        print("the periods")
        check(period_for("monthly", date(2026, 9, 5)) == (date(2026, 8, 1), date(2026, 8, 31)),
              "a monthly run on 5 September covers August")
        check(period_for("monthly", date(2026, 1, 28)) == (date(2025, 12, 1), date(2025, 12, 31)),
              "a monthly run on 28 January covers the December before")
        check(period_for("weekly", date(2026, 8, 17)) == (date(2026, 8, 10), date(2026, 8, 16)),
              "a Monday run covers the Monday to Sunday before")
        check(period_for("weekly", date(2026, 8, 23)) == (date(2026, 8, 10), date(2026, 8, 16)),
              "a Sunday run covers the week before its own")

        print("the runner")
        wid, mid = weekly["id"], monthly["id"]
        check(run_due(at("2026-09-04")) == [], "on a Friday neither schedule runs")
        last, exports, notices = state(mid)
        check(last is None and exports == [] and notices == [], "and nothing is written")

        unstored = run_due(at("2026-09-05"))
        check([(r["schedule_id"], r["status"]) for r in unstored] == [(mid, "storage_unconfigured")],
              f"on 5 September the monthly one is due, and storage is not configured ({unstored})")
        last, exports, notices = state(mid)
        check(last is None, "the schedule is not marked as run")
        check(exports == [] and notices == [], "no export and no notice is left behind")

        os.environ["S3_BUCKET"] = "mse-sched-check"
        built = run_due(at("2026-09-05"))
        check([(r["schedule_id"], r["status"]) for r in built] == [(mid, "built")],
              f"with storage, the monthly one is built and the weekly one is not ({built})")
        last, exports, notices = state(mid)
        check(last == at("2026-09-05"), "last_run_at is set to the run")
        check(len(exports) == 1 and exports[0][1] == "ready" and exports[0][5:7] == (date(2026, 8, 1), date(2026, 8, 31))
              and exports[0][2:5] == ("month_summary", "xlsx", "cash"),
              f"one ready export of August, as the schedule says ({exports[0][1:7] if exports else None})")
        check(len(notices) == 1 and notices[0][0] == exports[0][0], "one notice, naming the export")
        if notices:
            print(f"       notice: {notices[0][1]} {notices[0][2]}")

        again = run_due(at("2026-09-05") .replace(hour=20))
        check([(r["schedule_id"], r["status"]) for r in again] == [(mid, "already_run")],
              "a second run the same London day finds it already run")
        last2, exports2, notices2 = state(mid)
        check(last2 == last and len(exports2) == 1 and len(notices2) == 1, "and writes nothing")

        weekly_run = run_due(at("2026-08-17"))
        check([(r["schedule_id"], r["status"]) for r in weekly_run] == [(wid, "built")],
              f"on Monday 17 August the weekly one is built and the monthly one is not ({weekly_run})")
        _, exports, notices = state(wid)
        check(len(exports) == 2 and exports[1][5:7] == (date(2026, 8, 10), date(2026, 8, 16)),
              "the weekly export covers 10 to 16 August")
        check(len(notices) == 2, "one more notice, one per run")

        update_export_schedule(ScheduleChange(active=False), me, SHOP, uuid.UUID(wid))
        check(run_due(at("2026-08-24")) == [], "a paused weekly schedule does not run on its day")

        print("the files against the Money screen")
        for _export_id, _status, kind, fmt, basis, start, end, key in exports:
            rows = read_file(storage.get_bytes(key), fmt)
            got = totals_from_rows(rows)
            screen = get_money(response=Response(), account=me, shop_id=SHOP, basis=basis,
                               period_from=start, period_to=end, granularity="month", if_none_match=None)
            want = {
                "Gross sales": screen.totals.gross_sales.amount_minor,
                "Net sales": screen.totals.net_sales.amount_minor,
                "Net proceeds": screen.totals.net_proceeds.amount_minor,
            }
            if screen.kept is not None:
                want["Gross profit after returns"] = screen.kept.amount_minor
            label = f"{kind} {fmt} {basis} {start} to {end}"
            check(want["Gross sales"] != 0, f"{label}: the period holds sales ({want['Gross sales']})")
            check(all(got.get(k) == v for k, v in want.items()),
                  f"{label}: file totals equal the screen ({ {k: got.get(k) for k in want} } against {want})")
            check(rows[1][1] == f"{start.isoformat()} to {end.isoformat()}", f"{label}: the file names its period")
    finally:
        with pool().connection() as conn, conn.transaction():
            conn.execute("delete from notifications where account_id = %s and type = 'scheduled_export_ready'",
                         (str(ACCOUNT),))
            conn.execute("delete from exports where shop_id = %s", (str(SHOP),))
            conn.execute("delete from export_schedules where shop_id in (%s, %s)", (str(SHOP), str(other_shop)))
            conn.execute("delete from idempotency_keys where account_id in (%s, %s)", (str(ACCOUNT), str(other)))
            conn.execute("delete from shops where id = %s", (str(other_shop),))
            conn.execute("delete from accounts where id = %s", (str(other),))
        server.terminate()
    print(f"\n{ok} passed, {bad} failed")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
